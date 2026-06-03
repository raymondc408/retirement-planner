#!/usr/bin/env python3
"""
提早退休規劃計算器 - Streamlit 互動版
完全參數化 + 年度開支表 (可隨時改特定年份金額) + 多退休年份比較 + 所需本金反推 + 現有本金正向走勢模擬 (2-3 回報率比較) + 資本走勢表 (像你圖片的格式) + 完整 Excel 匯出
"""

import streamlit as st
import pandas as pd
import numpy as np
from typing import Dict, Any, List, Tuple
import datetime
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.chart import BarChart, LineChart, Reference
import plotly.express as px

st.set_page_config(
    page_title="Early Retirement Planner",
    page_icon="💰",
    layout="wide",
    initial_sidebar_state="expanded"
)

# ==================== 核心計算函式 (從之前的 Excel 版移植，改為年度表驅動) ====================

def simulate(
    initial_capital: float,
    annual_return: float,
    retirement_year: int,
    end_year: int,
    spending_dict: Dict[int, int],
    final_buffer_years: float,
) -> Dict[str, Any]:
    """使用年度開支表進行單一情境模擬。"""
    balances: List[float] = []
    withdrawals: List[int] = []
    years: List[int] = []

    portfolio = float(initial_capital)

    for y in range(retirement_year, end_year + 1):
        withdrawal = int(spending_dict.get(y, 0))
        portfolio -= withdrawal
        balances.append(portfolio)
        withdrawals.append(withdrawal)
        years.append(y)

        portfolio *= (1 + annual_return)

        if portfolio < 0:
            break

    final_balance = balances[-1] if balances else 0.0
    min_balance = min(balances) if balances else 0.0
    last_expense = withdrawals[-1] if withdrawals else 0
    required_buffer = last_expense * final_buffer_years

    success = (min_balance >= 0) and (final_balance >= required_buffer)

    return {
        "success": success,
        "min_balance": min_balance,
        "final_balance": final_balance,
        "balances": balances,
        "withdrawals": withdrawals,
        "years": years,
    }


def find_required_capital(
    annual_return: float,
    retirement_year: int,
    end_year: int,
    spending_dict: Dict[int, int],
    final_buffer_years: float,
    low: float = 1_000_000,
    high: float = 100_000_000,
    tol: float = 50_000,
) -> Tuple[float, Dict[str, Any]]:
    """二分搜尋找出退休當年所需本金 (使用年度開支表)。"""
    best_capital = high
    best_result: Dict[str, Any] = {}

    for _ in range(60):
        mid = (low + high) / 2
        result = simulate(mid, annual_return, retirement_year, end_year, spending_dict, final_buffer_years)

        if result["success"]:
            high = mid
            best_capital = mid
            best_result = result
        else:
            low = mid

        if high - low < tol:
            break

    final_result = simulate(best_capital, annual_return, retirement_year, end_year, spending_dict, final_buffer_years)
    return best_capital, final_result


def build_capital_trajectory(
    starting_capital: float,
    start_year: int,
    end_year: int,
    spending_dict: Dict[int, int],
    rate: float,
    base_current_year: int = 2026,
    base_current_age: int = 49,
) -> pd.DataFrame:
    """產生資本走勢表 (Simple Model 格式，像你圖片: Capital + Passive=cap*rate + Expense + Gain/Loss)。
    可用於退休當年所需本金的走勢，或從現有本金往前模擬。
    """
    rows = []
    capital = starting_capital

    for idx, y in enumerate(range(start_year, end_year + 1)):
        age = base_current_age + (y - base_current_year)
        expense = spending_dict.get(y, 0)
        passive = capital * rate
        gain_loss = passive - expense
        next_capital = capital + gain_loss

        rows.append({
            "No.": idx + 1,
            "Age": age,
            "YEAR": y,
            f"Capital (HKD) {rate*100:.1f}%": round(capital),
            f"Passive Income (HKD) {rate*100:.1f}%": round(passive),
            "Expense (HKD)": expense,
            f"Gain / Loss (HKD) {rate*100:.1f}%": round(gain_loss),
        })
        capital = next_capital

    return pd.DataFrame(rows)

# 向後相容舊名稱
def build_simple_model_table(*args, **kwargs):
    return build_capital_trajectory(*args, **kwargs)


# ==================== Streamlit UI ====================

# Language selector at top (default English as requested)
lang = st.selectbox(
    "Language / 語言",
    ["English", "中文繁體"],
    index=0,
    key="lang_selector"
)
lang_code = "en" if lang == "English" else "zh"

def _t(key):
    """Translation helper. All UI strings centralized here for bilingual support (en/zh).
    Named _t (not bare _) to avoid being shadowed by loop variables like `for _, row in df.iterrows()`.
    """
    texts = {
        "title": {
            "en": "💰 Early Retirement Planner - Streamlit Interactive Version",
            "zh": "💰 提早退休規劃計算器 - Streamlit 互動版"
        },
        "caption": {
            "en": "Fully parameterized yearly spending table + multi-retirement year comparison + capital trajectory table (matching your uploaded image format)",
            "zh": "完全參數化年度開支表 + 多退休年份比較 + 資本走勢表 (類似你上傳圖片的格式)"
        },
        "usage": {
            "en": """**Usage**:
1. Edit the **Yearly Spending Table** on the left or below (directly change amounts for any year. All item names (Living Cost / Education Items / Mortgage / Other) can be customized in the sidebar).
2. Set basic parameters and the retirement years + return rates you want to compare (multiple selectable).
3. Click "Recalculate" or the tool will update automatically.
4. The upper area shows "Required Capital Comparison" (back-calculate how much to prepare for retirement) + charts + Simple Model trajectory tables.
5. Below is the new **Existing Capital Trajectory Simulation**: input your current capital, select 2-3 return rates, and instantly calculate and plot charts + detailed tables using the same spending table.
6. Click "Generate Excel Report" to export **complete data**.""",
            "zh": """**使用方式**：
1. 在左側或下方編輯 **年度開支表**（直接改任何年份的各項目金額。所有項目名稱（Living Cost / 教育項目 / Mortgage / Other）可在側邊欄自訂）。
2. 設定基本參數與想比較的退休年份 + 回報率（可多選）。
3. 點擊「重新計算」或工具會自動更新。
4. 上方區域可看「所需本金比較」（反推退休當年要準備多少）+ 圖表 + Simple Model 走勢表。
5. 下方新增 **現有本金走勢模擬**：自行輸入目前本金，挑 2-3 個回報率，即時用同一開支表計算並畫圖表 + 詳細表。
6. 按「產生 Excel 報告」會匯出**完整所有資料**。"""
        },
        # Sidebar
        "basic_settings": {"en": "Basic Settings", "zh": "基本設定"},
        "current_age": {"en": "Current Age", "zh": "目前年齡"},
        "current_year": {"en": "Current Year", "zh": "目前年份"},
        "end_age": {"en": "Plan to Age (End Age)", "zh": "規劃活到年齡"},
        "buffer_years": {"en": "Final Buffer Years (capital must remain >=0 + last year expense * buffer)", "zh": "末年緩衝年數（資本須維持 >=0 + 最後一年開支 × 緩衝）"},
        "passive_return_rates": {"en": "📈 Passive Return Rates", "zh": "📈 被動收入回報率"},
        "select_rates": {"en": "Select rates to compare (multi-select 1%–6%)", "zh": "選擇要比較的回報率（可多選，1%–6%）"},
        "retirement_scenarios": {"en": "Retirement Scenarios (Year(s))", "zh": "退休情境年份列表"},
        "retirement_years_list": {"en": "Retirement year(s), comma separated (e.g. 2028 or 2028,2032)", "zh": "退休年份列表（逗號分隔，例如 2028 或 2028,2032）"},
        "retirement_caption": {"en": "To see retirement at age 51 (2028), use 2028, or enter 2028,2032 to compare multiple scenarios", "zh": "想看 51歲退休 (2028) 就用 2028，或輸入 2028,2032 比較多個情境"},
        "custom_expense_names": {"en": "Customize Expense Item Display Names", "zh": "自訂開支項目顯示名稱"},
        "living_name": {"en": "Living Cost name", "zh": "生活開支名稱"},
        "edu1_name": {"en": "1st Education item name", "zh": "第一教育項目名稱"},
        "edu2_name": {"en": "2nd Education item name", "zh": "第二教育項目名稱"},
        "mort_name": {"en": "Mortgage name", "zh": "按揭名稱"},
        "other_name": {"en": "Other name", "zh": "其他開支名稱"},
        "name_change_note": {"en": "Changes apply instantly to table columns, preview, calculations & export. These are your custom display names (not auto-translated when you switch language).", "zh": "修改會即時套用至表格欄位、預覽、所有計算與 Excel 匯出。這些是你自訂的顯示名稱（切換語言時不會自動翻譯）。"},
        # Zone 1
        "spending_table_header": {"en": "Yearly Spending Table (Editable - single source of truth)", "zh": "年度開支表（可編輯 - 所有計算的唯一來源）"},
        "spending_zone_desc": {"en": "Directly edit amounts for any specific year. Customize the 5 item names in the sidebar (Living / Kate / Damon / Mortgage / Other). The formatted preview below shows live values without editing cells.", "zh": "直接修改任何特定年份的金額。5 個項目名稱可在側邊欄自訂（生活費 / Kate / Damon / 按揭 / 其他）。下方格式化預覽可即時查看數值，無需進入編輯格。"},
        "edit_expander": {"en": "✏️ Click to expand/collapse editing table (modify values)", "zh": "✏️ 點擊展開/收起 編輯表格（修改數值）"},
        "preview_header": {"en": "📋 Formatted Preview (live view of current values, no clicks needed)", "zh": "📋 格式化預覽（即時查看目前各年實際數值，無需點擊編輯）"},
        "preview_caption": {"en": "This preview is formatted for easy reading (HKD x,xxx,xxx). It always reflects the latest edits from the table above. To change any amount, expand the editor and modify directly.", "zh": "此預覽已格式化方便閱讀（HKD x,xxx,xxx）。永遠反映上方表格的最新編輯。要更改金額，請展開編輯區直接修改對應欄位。"},
        "recalculate_button": {"en": "🔄 Recalculate / Refresh Projections", "zh": "🔄 重新計算 / 刷新所有預測"},
        # Zone 2
        "required_capital_header": {"en": "Required Capital Comparison (across retirement years & rates)", "zh": "所需本金比較（不同退休年份與回報率）"},
        "required_zone_desc": {"en": "Binary search finds the starting capital needed at each retirement year so your portfolio lasts to end age with buffer. All driven by the spending table.", "zh": "使用二分搜尋計算各退休年份開始時需準備的本金，使組合能維持至規劃末年（含緩衝）。完全由上方年度開支表驅動。"},
        "main_retire_label": {"en": "Select main retirement year to view detailed charts & Simple Model tables", "zh": "選擇主要退休年份以查看詳細圖表與 Simple Model 走勢表"},
        "capital_chart_title": {"en": "Capital Trajectory (Compound Model) - Retirement Year {year}", "zh": "資本餘額走勢圖（複利模型）- 退休年 {year}"},
        "simple_model_subheader": {"en": "📋 Capital Trajectory Tables (Simple Model - exact format from your images)", "zh": "📋 資本走勢表 (Simple Model - 完全依照你圖片格式)"},
        "simple_model_caption": {"en": "Simple Model: each year Passive Income = Capital at start of year × rate, minus Expense = Gain/Loss, added directly to capital. Expenses come 100% from the yearly spending table you edited.", "zh": "此表使用簡單模型：每年 Passive Income = 當年初 Capital × 利率，減掉該年 Expense = Gain/Loss，直接加減到資本。開支完全來自上面你編輯的年度開支表。"},
        "no_rate_for_table": {"en": "Please select at least one return rate in the sidebar to display the capital trajectory tables.", "zh": "請在側邊欄選擇至少一個回報率以顯示資本走勢表。"},
        # Zone 3
        "forward_header": {"en": "Existing Capital Trajectory Simulation (Forward Projection)", "zh": "現有本金走勢模擬（正向預測）"},
        "forward_zone_desc": {"en": "Input your actual present capital, pick 2-3 rates from above; instantly projects year-by-year using IDENTICAL spending table + retire year from Basic Settings. Includes sensitivity analysis.", "zh": "輸入你目前的實際本金，從上方挑 2-3 個回報率；即時使用與「基本設定」完全相同的開支表 + 退休年期做逐年正向模擬。包含敏感度分析。"},
        "forward_caption": {"en": "Start year auto-syncs to the earliest retirement scenario year you listed. Living costs & all expenses pulled live from the single editable yearly spending table.", "zh": "開始年期自動對齊你側邊欄「退休情境年份列表」的最早年份。生活成本與所有開支完全來自同一個可編輯年度開支表。"},
        "current_capital_label": {"en": "Your current / starting capital for forward sim (HKD)", "zh": "你目前的起始本金（用於正向模擬，HKD）"},
        "current_capital_help": {"en": "Used ONLY for the forward simulation + sensitivity below. Does NOT change the Required Capital calculations in Zone 2.", "zh": "僅用於下方的正向模擬與敏感度分析。不會影響第 2 區的「所需本金」計算。"},
        "forward_rate_pick": {"en": "From the multi-select above \"📈 Passive Return Rates\", pick the ones to compare this time (recommend 2-3)", "zh": "從上方「📈 被動收入回報率」多選中挑選要比較的（建議 2-3 個）"},
        "fwd_start_caption": {"en": "✅ Simulation start year auto-aligned to basic settings: **{year}** (from sidebar retirement years list). Spending fully from the same editable table.", "zh": "✅ 模擬開始年期自動對齊基本設定：**{year}**（來自側邊欄「退休情境年份列表」）。生活成本/開支完全來自同一個可編輯年度開支表。"},
        "fwd_chart_title": {"en": "Existing Capital Projection Comparison (start capital HKD {capital:,.0f}, from retire year {year})", "zh": "現有本金走勢比較 (起始本金 HKD {capital:,.0f}，從退休年 {year} 開始)"},
        "fwd_detail_caption": {"en": "⚠️ This sim's start year is locked to your basic retirement scenario settings. Expenses = exact numbers from the spending editor.", "zh": "⚠️ 本模擬的開始年期已自動與基本設定一致（使用側邊欄「退休情境年份列表」的最早年份 {year}），生活成本/開支完全來自同一個可編輯年度開支表。"},
        "fwd_detail_header": {"en": "**Detailed Yearly Tables (Simple Model format)**", "zh": "**詳細每年走勢表 (Simple Model 格式)**"},
        "select_at_least_one_rate": {"en": "Please select at least 1 return rate in the sidebar or here to see charts and tables.", "zh": "請在上方側邊欄或此處選擇至少 1 個回報率，即可看到本金走勢圖表與表格。"},
        "select_rates_for_sim": {"en": "Select 2-3 rates for this forward sim (from the ones chosen above)", "zh": "為本次正向模擬選擇 2-3 個回報率（從上方已選的）"},
        # Sensitivity
        "sensitivity_header": {"en": "💡 Starting Capital Sensitivity Analysis", "zh": "💡 起始資本敏感度分析"},
        "sensitivity_caption": {"en": "Impact of base / +1M / +2M starting capital on ending capital at plan's end year (for the rates selected in the forward simulation). Compounding amplifies differences over long periods.", "zh": "基本 / 多加 100萬 / 多加 200萬 起始本金對計劃末年剩餘本金的影響（使用正向模擬所選回報率）。長時間複利會放大差異。"},
        "sensitivity_no_data": {"en": "Select 1+ forward rates above to compute sensitivity table & chart.", "zh": "請在上方選擇回報率以計算敏感度分析表與圖表。"},
        "main_info": {"en": "💡 Observation: Adding HKD 1M–2M today can create ~4M difference in terminal wealth after decades of compounding (depends on rates & spending path).", "zh": "💡 觀察：多準備 100–200 萬，經數十年複利後，晚年餘額差距可能達到約 400 萬（視回報率與開支路徑而定）。"},
        "sensitivity_reminder": {"en": "All sensitivity runs use the exact same spending_dict + end_year + rates as the forward projection section (which itself mirrors your Basic Settings + spending table).", "zh": "所有敏感度計算都使用與正向模擬區完全相同的 spending 數字、結束年份與回報率（而正向模擬本身與基本設定 + 開支表一致）。"},
        # Export
        "export_button": {"en": "📤 Generate Excel Report (contains ALL current inputs + every calculation)", "zh": "📤 產生 Excel 報告（包含你目前所有輸入 + 所有計算結果）"},
        "export_summary_title": {"en": "Retirement Planning Summary & Required Capital Matrix", "zh": "退休規劃摘要與所需本金矩陣"},
        "basic_params": {"en": "Basic Parameters", "zh": "基本參數"},
        "current_age_label": {"en": "Current Age", "zh": "目前年齡"},
        "current_year_label": {"en": "Current Year", "zh": "目前年份"},
        "end_age_label": {"en": "End Age (plan to live to)", "zh": "規劃活到年齡"},
        "buffer_label": {"en": "Final Buffer (years)", "zh": "末年緩衝年數"},
        "selected_rates_label": {"en": "Selected Passive Return Rates", "zh": "已選被動回報率"},
        "retirement_years_label": {"en": "Retirement Year(s) in Comparison", "zh": "比較中的退休年份"},
        "required_capital_title": {"en": "Required Starting Capital at Retirement Year", "zh": "退休當年所需起始本金"},
        "retirement_year_col": {"en": "Retirement Year", "zh": "退休年份"},
        "note_projections": {"en": "Full year-by-year details for every retirement scenario are in the Projections_YYYY and Simple_Model_YYYY sheets.", "zh": "每個退休情境的完整逐年細節請見 Projections_YYYY 與 Simple_Model_YYYY 工作表。"},
        "full_spending_title": {"en": "Full Yearly Spending Table (all years, live from editor)", "zh": "完整年度開支表（所有年份，直接來自編輯器）"},
        "export_time": {"en": "Exported at", "zh": "匯出時間"},
        "full_spending_note": {"en": "Exact spending numbers used for ALL projections & required-capital searches. Item column headers reflect your custom names from the sidebar.", "zh": "所有預測與所需本金計算實際使用的開支數字。項目欄位標題會反映你在側邊欄自訂的名稱。"},
        "download_label": {"en": "⬇️ Download Full Excel Report", "zh": "⬇️ 下載完整 Excel 報告"},
        "export_filename": {"en": "Retirement_Plan", "zh": "退休規劃報告"},
        "disclaimer": {"en": "Disclaimer: This is a simplified deterministic model for educational/planning purposes only. Real returns vary, inflation (removed per your request), taxes, longevity, healthcare costs etc. are not modeled. Not financial advice. Consult qualified professionals.", "zh": "免責聲明：這是簡化確定性模型，僅供教育與規劃參考。你已要求移除通脹。實際回報有波動，稅務、醫療、長壽風險等未納入。非財務建議，請諮詢合格專業人士。"},
        # Misc / forward other
        "summary_ret_year": {"en": "Retirement Year", "zh": "退休年份"},
        "summary_age": {"en": "Approx. Age", "zh": "對應年齡(約)"},
        "summary_req_cap": {"en": "{rate}% Required Capital", "zh": "{rate}% 所需本金"},
        "export_subheader": {"en": "Export Current Scenario to Excel", "zh": "匯出目前情境成 Excel"},
        "tip_caption": {"en": "Tip: The most powerful editing experience is to directly modify in the \"Yearly Spending Table\" above in real time + recalculate. If you want to save for long-term use, you can keep the parameters.xlsx as well (the Excel version generator is still available).", "zh": "提示：最強大的編輯體驗還是直接在上面的「年度開支表」即時修改 + 重算。想存檔長期使用，可以把 parameters.xlsx 也一起保留 (Excel 版產生器還是可用的)。"},
    }
    val = texts.get(key, {}).get(lang_code, key)
    return val

st.title(_t("title"))
st.caption(_t("caption"))

st.markdown(_t( "usage" ))

# --- Sidebar 基本設定 ---
with st.sidebar:
    st.header(_t("basic_settings"))

    current_age = st.number_input(_t("current_age"), min_value=35, max_value=70, value=49, step=1)
    current_year = st.number_input(_t("current_year"), min_value=2020, max_value=2035, value=2026, step=1)
    end_age = st.number_input(_t("end_age"), min_value=80, max_value=110, value=100, step=1)
    buffer_years = st.slider(_t("buffer_years"), 0.0, 3.0, 1.5, 0.5)

    st.divider()
    st.header(_t("passive_return_rates"))
    # 支援從 1% 到 6%，包含保守到較進取的選項
    rate_options = [0.01, 0.015, 0.02, 0.025, 0.03, 0.035, 0.04, 0.045, 0.05, 0.055, 0.06]
    selected_rates = st.multiselect(
        _t("select_rates"),
        rate_options,
        default=[0.03, 0.035, 0.04],
        format_func=lambda x: f"{x*100:.1f}%"
    )

    st.divider()
    st.header(_t("retirement_scenarios"))
    ret_input = st.text_input(_t("retirement_years_list"), value="2028")
    try:
        retirement_years = sorted([int(x.strip()) for x in ret_input.split(",") if x.strip()])
    except:
        retirement_years = [2028]
    if not retirement_years:
        retirement_years = [2028]

    st.caption(_t( "retirement_caption" ))

    st.divider()
    st.header(_t("custom_expense_names"))
    living_name = (st.text_input(_t("living_name"), value="Living Cost") or "Living Cost").strip()
    edu1_name = (st.text_input(_t("edu1_name"), value="Kate 學費") or "Kate 學費").strip()
    edu2_name = (st.text_input(_t("edu2_name"), value="Damon 學費") or "Damon 學費").strip()
    mort_name = (st.text_input(_t("mort_name"), value="Mortgage") or "Mortgage").strip()
    other_name = (st.text_input(_t("other_name"), value="Other") or "Other").strip()
    st.caption(_t("name_change_note"))

# --- Zone 1: 年度開支表 ---
st.markdown(f"""
<div style="background-color: #E3F2FD; padding: 14px; border-radius: 10px; border-left: 8px solid #1565C0; margin: 10px 0 15px 0;">
    <h3 style="color: #0D47A1; margin: 0; display: flex; align-items: center;">
        📅 1. {_t("spending_table_header")}
    </h3>
    <p style="color: #1565C0; font-size: 0.9em; margin: 5px 0 0 0;">{_t("spending_zone_desc")}</p>
</div>
""", unsafe_allow_html=True)

# 動態欄位名稱（來自側邊欄自訂）
living_col = f"{living_name} (HKD)"
edu1_col = f"{edu1_name} (HKD)"
edu2_col = f"{edu2_name} (HKD)"
mort_col = f"{mort_name} (HKD)"
other_col = f"{other_name} (HKD)"
# For display in texts and preview (base name without (HKD))
living_display = living_name
edu1_display = edu1_name
edu2_display = edu2_name
mort_display = mort_name
other_display = other_name

st.info(
    _t("preview_caption") + f"  |  " + 
    (f"Edit amounts via the expander above. Columns: {living_display} / {edu1_display} / {edu2_display} / {mort_display} / {other_display}." if lang_code == "en" else f"請展開上方編輯器修改金額。欄位：{living_display} / {edu1_display} / {edu2_display} / {mort_display} / {other_display}。")
)

# 動態年份範圍：涵蓋從 2025 年到根據目前年齡、目前年份、規劃活到年齡計算的結束年
# 這樣如果用戶把活到年齡設超過 100 歲，表格會自動延伸，並在 2039 年之後繼續預設每年 Living Cost 600,000
planning_end_year = current_year + (end_age - current_age)
year_range = list(range(2025, planning_end_year + 1))

# 建立初始 DataFrame - 完全依照你最新的更正
# 2025-2027: 全 0
# 2028-2031: Living Cost 600000, 第一教育項目 300000, 第二教育項目 300000, Mortgage 0, Other 0
# 2032-2034: Living Cost 600000, 第一教育項目 500000, 第二教育項目 300000, Mortgage 0, Other 0
# 2035: Living Cost 600000, 第一教育項目 0, 第二教育項目 300000, Mortgage 0, Other 0
# 2036-2038: Living Cost 600000, 第一教育項目 0, 第二教育項目 500000, Mortgage 0, Other 0
# 2039 之後（直到規劃結束年）：Living Cost 600000, 第一/第二教育項目 0, Mortgage 0, Other 0

living = []
edu1 = []  # 第一個自訂教育開支項目
edu2 = []  # 第二個自訂教育開支項目
mort = []
other = []

for y in year_range:
    if 2025 <= y <= 2027:
        living.append(0)
        edu1.append(0)
        edu2.append(0)
        mort.append(0)
        other.append(0)
    elif 2028 <= y <= 2031:
        living.append(600000)
        edu1.append(300000)
        edu2.append(300000)
        mort.append(0)
        other.append(0)
    elif 2032 <= y <= 2034:
        living.append(600000)
        edu1.append(500000)
        edu2.append(300000)
        mort.append(0)
        other.append(0)
    elif y == 2035:
        living.append(600000)
        edu1.append(0)
        edu2.append(300000)
        mort.append(0)
        other.append(0)
    elif 2036 <= y <= 2038:
        living.append(600000)
        edu1.append(0)
        edu2.append(500000)
        mort.append(0)
        other.append(0)
    else:  # 2039 到 2080
        living.append(600000)
        edu1.append(0)
        edu2.append(0)
        mort.append(0)
        other.append(0)

data = {
    "Year": year_range,
    living_col: living,
    edu1_col: edu1,
    edu2_col: edu2,
    mort_col: mort,
    other_col: other,
}
spending_df = pd.DataFrame(data).astype({
    living_col: "int64",
    edu1_col: "int64",
    edu2_col: "int64",
    mort_col: "int64",
    other_col: "int64",
})

# 可編輯表格放在 expander 內，預設收起，節省空間
with st.expander(_t( "edit_expander" ), expanded=False):
    edited_spending = st.data_editor(
        spending_df,
        use_container_width=True,
        num_rows="fixed",
        key="yearly_spending",
        column_config={
            "Year": st.column_config.NumberColumn(disabled=True),
            living_col: st.column_config.NumberColumn(step=10000),
            edu1_col: st.column_config.NumberColumn(step=10000),
            edu2_col: st.column_config.NumberColumn(step=10000),
            mort_col: st.column_config.NumberColumn(step=10000),
            other_col: st.column_config.NumberColumn(step=10000),
        }
    )

# 格式化預覽（方便一眼看到所有數值，無需點擊每個儲存格） - 使用最新編輯的值
st.markdown("**" + _t( "preview_header" ) + "**")
preview_df = edited_spending.copy()
# Format values using the actual column names present in the DataFrame
for col in [living_col, edu1_col, edu2_col, mort_col, other_col]:
    if col in preview_df.columns:
        preview_df[col] = preview_df[col].apply(lambda x: f"HKD {int(x):,}" if int(x) != 0 else "HKD 0")
# Short display names for preview (always with (HKD) for consistency)
preview_df = preview_df.rename(columns={
    living_col: f"{living_display} (HKD)",
    edu1_col: f"{edu1_display} (HKD)",
    edu2_col: f"{edu2_display} (HKD)",
    mort_col: f"{mort_display} (HKD)",
    other_col: f"{other_display} (HKD)",
})
st.dataframe(preview_df, use_container_width=True, hide_index=True)

st.caption(_t("preview_caption"))

# 建立 spending_dict （從最新編輯的表格產生，用於後續所有計算）
spending_dict = {}
for idx, row in edited_spending.iterrows():
    total = int(row[living_col]) + int(row[edu1_col]) + \
            int(row[edu2_col]) + int(row[mort_col]) + int(row[other_col])
    spending_dict[int(row["Year"])] = total

# 計算按鈕
if st.button(_t("recalculate_button"), type="primary", use_container_width=True):
    st.session_state["last_calc"] = datetime.datetime.now()

# 自動計算 (第一次或有變動)
if "last_calc" not in st.session_state:
    st.session_state["last_calc"] = datetime.datetime.now()

# ==================== 計算 ====================
end_year = current_year + (end_age - current_age)

results = {}  # retirement_year -> {rate: required_capital}

for ret_y in retirement_years:
    results[ret_y] = {}
    for r in selected_rates:
        req, _result = find_required_capital(
            annual_return=r,
            retirement_year=ret_y,
            end_year=end_year,
            spending_dict=spending_dict,
            final_buffer_years=buffer_years,
        )
        results[ret_y][r] = req

# ==================== 輸出區 ====================

st.divider()
# --- Zone 2: 所需本金比較 ---
st.markdown(f"""
<div style="background-color: #E8F5E9; padding: 14px; border-radius: 10px; border-left: 8px solid #2E7D32; margin: 10px 0 15px 0;">
    <h3 style="color: #1B5E20; margin: 0; display: flex; align-items: center;">
        📈 2. {_t("required_capital_header")}
    </h3>
    <p style="color: #2E7D32; font-size: 0.9em; margin: 5px 0 0 0;">{_t("required_zone_desc")}</p>
</div>
""", unsafe_allow_html=True)

if results:
    summary_rows = []
    for ret_y in retirement_years:
        age = current_age + (ret_y - current_year)
        row = {_t( "summary_ret_year" ): ret_y, _t( "summary_age" ): age}
        for r in selected_rates:
            pct = f"{r*100:.1f}%"
            # Build header using template (replace {rate} after)
            hdr = _t( "summary_req_cap" ).format(rate=pct)
            row[hdr] = results[ret_y][r]
        summary_rows.append(row)

    summary_df = pd.DataFrame(summary_rows)
    st.dataframe(summary_df.style.format({col: "HKD {:,.0f}" for col in summary_df.columns if "%" in col or "Required" in col or "所需" in col}), use_container_width=True)

    # 選擇主要退休年份看詳細
    main_ret = st.selectbox(_t( "main_retire_label" ), retirement_years, index=0)

    st.divider()
    st.subheader( _t( "capital_chart_title" ).format(year=main_ret) )

    # 為每個 rate 跑 simulate 取得 balances
    chart_rows = []
    for r in selected_rates:
        req_cap = results[main_ret][r]
        sim = simulate(req_cap, r, main_ret, end_year, spending_dict, buffer_years)
        for i, bal in enumerate(sim["balances"]):
            chart_rows.append({
                "Year": sim["years"][i],
                "Capital": bal,
                "Rate": f"{r*100:.1f}%"
            })

    if chart_rows:
        chart_df = pd.DataFrame(chart_rows)
        import plotly.express as px
        fig = px.line(
            chart_df, x="Year", y="Capital", color="Rate",
            title=_t( "capital_chart_title" ).format(year=main_ret),
            labels={"Capital": "Capital (HKD)" if lang_code=="en" else "資本餘額 (HKD)", "Year": "Year" if lang_code=="en" else "年份"}
        )
        fig.update_layout(height=450)
        st.plotly_chart(fig, use_container_width=True)

    st.divider()
    st.subheader(_t( "simple_model_subheader" ))

    st.caption(_t( "simple_model_caption" ))

    # 取前兩個 (或全部如果少於兩個) 來顯示簡單模型表 (像你圖片的格式)。若無選擇則跳過。
    display_rates = selected_rates[:2] if selected_rates else []

    simple_dfs = []
    for r in display_rates:
        if r not in results.get(main_ret, {}):
            continue
        req = results[main_ret][r]
        df_simple = build_simple_model_table(
            starting_capital=req,
            start_year=main_ret,
            end_year=end_year,
            spending_dict=spending_dict,
            rate=r,
            base_current_year=current_year,
            base_current_age=current_age
        )
        # 重新命名欄位讓它更像圖片，並確保有 HKD
        df_simple = df_simple.rename(columns={
            f"Capital (HKD) {r*100:.1f}%": "Capital (HKD)",
            f"Passive Income (HKD) {r*100:.1f}%": "Passive Income (HKD)",
            "Expense (HKD)": "Expense (HKD)",
            f"Gain / Loss (HKD) {r*100:.1f}%": "Gain / Loss (HKD)",
        })
        df_simple["Rate"] = f"{r*100:.1f}%"
        simple_dfs.append(df_simple)

    # 合併顯示 (或分開兩個表格)
    if len(simple_dfs) >= 2:
        col1, col2 = st.columns(2)
        with col1:
            st.markdown(f"**{display_rates[0]*100:.1f}%**")
            disp1 = simple_dfs[0][["No.", "Age", "YEAR", "Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]].copy()
            for c in ["Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]:
                disp1[c] = disp1[c].apply(lambda x: f"HKD {x:,.0f}")
            st.dataframe(disp1, use_container_width=True, hide_index=True)
        with col2:
            st.markdown(f"**{display_rates[1]*100:.1f}%**")
            disp2 = simple_dfs[1][["No.", "Age", "YEAR", "Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]].copy()
            for c in ["Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]:
                disp2[c] = disp2[c].apply(lambda x: f"HKD {x:,.0f}")
            st.dataframe(disp2, use_container_width=True, hide_index=True)
    elif len(simple_dfs) == 1:
        st.markdown(f"**{display_rates[0]*100:.1f}%**")
        disp0 = simple_dfs[0][["No.", "Age", "YEAR", "Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]].copy()
        for c in ["Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]:
            disp0[c] = disp0[c].apply(lambda x: f"HKD {x:,.0f}")
        st.dataframe(disp0, use_container_width=True, hide_index=True)
    else:
        st.caption(_t( "no_rate_for_table" ))

# ==================== 新功能：現有本金走勢模擬 (Forward Projection) ====================
st.divider()
# --- Zone 3: 現有本金走勢模擬 ---
st.markdown(f"""
<div style="background-color: #FFF3E0; padding: 14px; border-radius: 10px; border-left: 8px solid #E65100; margin: 10px 0 15px 0;">
    <h3 style="color: #BF360C; margin: 0; display: flex; align-items: center;">
        💰 3. {_t("forward_header")}
    </h3>
    <p style="color: #E65100; font-size: 0.9em; margin: 5px 0 0 0;">{_t("forward_zone_desc")}</p>
</div>
""", unsafe_allow_html=True)

st.caption(_t("forward_caption"))

# 輸入區
col_cap, col_rates = st.columns([1, 2])
with col_cap:
    current_capital = st.number_input(
        _t("current_capital_label"), 
        min_value=0, 
        value=15000000, 
        step=500000,
        help=_t("current_capital_help")
    )
with col_rates:
    st.write( _t( "forward_rate_pick" ) )
    if selected_rates:
        default_fwd = selected_rates[:min(3, len(selected_rates))]
        forward_rates = st.multiselect(
            _t("select_rates_for_sim"),
            options=selected_rates,
            default=default_fwd,
            format_func=lambda x: f"{x*100:.1f}%"
        )
    else:
        forward_rates = []

# 與基本設定一致的開始年期（提前計算，方便顯示）
if 'retirement_years' in locals() and retirement_years:
    fwd_start_year = min(retirement_years)
else:
    fwd_start_year = current_year

st.caption( _t( "fwd_start_caption" ).format(year=fwd_start_year) )

if forward_rates:
    # 準備 forward 資料（從退休年期開始，生活成本使用同一年度開支表）
    fwd_chart_rows = []
    fwd_dfs = []
    for r in forward_rates:
        df_fwd = build_capital_trajectory(
            starting_capital=float(current_capital),
            start_year=fwd_start_year,
            end_year=end_year,
            spending_dict=spending_dict,
            rate=r,
            base_current_year=current_year,
            base_current_age=current_age
        )
        # 為了顯示，重新命名為通用欄位（像 Simple Model 區）
        df_fwd = df_fwd.rename(columns={
            f"Capital (HKD) {r*100:.1f}%": "Capital (HKD)",
            f"Passive Income (HKD) {r*100:.1f}%": "Passive Income (HKD)",
            "Expense (HKD)": "Expense (HKD)",
            f"Gain / Loss (HKD) {r*100:.1f}%": "Gain / Loss (HKD)",
        })
        df_fwd["Rate"] = f"{r*100:.1f}%"
        fwd_dfs.append(df_fwd)

        # chart 用 numeric 資料
        for idx, row in df_fwd.iterrows():
            fwd_chart_rows.append({
                "Year": int(row["YEAR"]),
                "Capital": int(row["Capital (HKD)"]),
                "Rate": row["Rate"]
            })

    # 圖表
    if fwd_chart_rows:
        fwd_chart_df = pd.DataFrame(fwd_chart_rows)
        fig = px.line(
            fwd_chart_df, 
            x="Year", 
            y="Capital", 
            color="Rate",
            title=_t( "fwd_chart_title" ).format(capital=current_capital, year=fwd_start_year),
            labels={"Capital": "Capital (HKD)" if lang_code=="en" else "資本餘額 (HKD)", "Year": "Year" if lang_code=="en" else "年份"}
        )
        fig.update_layout(height=450)
        st.plotly_chart(fig, use_container_width=True)

    st.caption( _t( "fwd_detail_caption" ).format(year=fwd_start_year) )

    # 詳細表格 (2-3 欄並排)
    st.markdown( _t( "fwd_detail_header" ) )
    num_show = min(3, len(fwd_dfs))
    if num_show > 0:
        tbl_cols = st.columns(num_show)
        for i in range(num_show):
            with tbl_cols[i]:
                r = forward_rates[i]
                st.markdown(f"**{r*100:.1f}%**")
                disp = fwd_dfs[i][["No.", "Age", "YEAR", "Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]].copy()
                for c in ["Capital (HKD)", "Passive Income (HKD)", "Expense (HKD)", "Gain / Loss (HKD)"]:
                    disp[c] = disp[c].apply(lambda x: f"HKD {int(x):,}")
                st.dataframe(disp, use_container_width=True, hide_index=True)

    # ========== 起始資本敏感度分析 ==========
    st.markdown(f"**{_t('sensitivity_header')}**")
    st.caption(_t("sensitivity_caption"))

    # Use lang-aware column names for the sensitivity table
    sens_start_label = "Starting Capital (HKD)" if lang_code == "en" else "起始本金"
    sens_end_label_tmpl = "{rate} Ending Capital" if lang_code == "en" else "{rate} 晚年本金"
    sens_diff_label_tmpl = "{rate} Difference" if lang_code == "en" else "{rate} 差額"

    extra_amounts = [0, 1_000_000, 2_000_000]
    sensitivity_rows = []
    chart_data = []
    for extra in extra_amounts:
        start_cap = current_capital + extra
        row = {sens_start_label: f"HKD {start_cap:,.0f}"}
        for idx, r in enumerate(forward_rates[:num_show]):
            df = build_capital_trajectory(
                starting_capital=float(start_cap),
                start_year=fwd_start_year,
                end_year=end_year,
                spending_dict=spending_dict,
                rate=r,
                base_current_year=current_year,
                base_current_age=current_age
            )
            df = df.rename(columns={
                f"Capital (HKD) {r*100:.1f}%": "Capital (HKD)",
                f"Passive Income (HKD) {r*100:.1f}%": "Passive Income (HKD)",
                "Expense (HKD)": "Expense (HKD)",
                f"Gain / Loss (HKD) {r*100:.1f}%": "Gain / Loss (HKD)",
            })
            if not df.empty:
                final_cap = int(df.iloc[-1]["Capital (HKD)"])
                base_final = int(fwd_dfs[idx].iloc[-1]["Capital (HKD)"]) if idx < len(fwd_dfs) else 0
                diff = final_cap - base_final
                rate_pct = f"{r*100:.1f}%"
                row[sens_end_label_tmpl.format(rate=rate_pct)] = f"HKD {final_cap:,.0f}"
                row[sens_diff_label_tmpl.format(rate=rate_pct)] = f"HKD {diff:,.0f}"
                # For chart, always use neutral English keys + translated labels in fig
                chart_data.append({
                    "start_cap": start_cap,
                    "rate": rate_pct,
                    "end_cap": final_cap,
                    "end_cap_M": round(final_cap / 1_000_000, 1)
                })
        sensitivity_rows.append(row)

    if sensitivity_rows:
        sens_df = pd.DataFrame(sensitivity_rows)
        st.dataframe(sens_df, use_container_width=True, hide_index=True)

        if chart_data:
            chart_df = pd.DataFrame(chart_data)
            title_sens = "Impact of Starting Capital on Terminal Wealth by Rate" if lang_code == "en" else "不同起始本金在各回報率下的晚年預計剩餘本金"
            fig = px.bar(
                chart_df,
                x="start_cap",
                y="end_cap_M",
                color="rate",
                barmode="group",
                text="end_cap_M",
                title=title_sens,
                labels={
                    "start_cap": "Starting Capital (HKD)" if lang_code == "en" else "起始本金 (HKD)",
                    "end_cap_M": "Terminal Capital (Millions HKD)" if lang_code == "en" else "晚年本金 (百萬 HKD)",
                    "rate": "Return Rate" if lang_code == "en" else "回報率"
                }
            )
            fig.update_traces(texttemplate="%{text:.1f}M", textposition="outside")
            fig.update_layout(
                height=450,
                xaxis_tickformat="~s",
                yaxis_tickformat=".0f",
                yaxis_ticksuffix="M"
            )
            st.plotly_chart(fig, use_container_width=True)
    else:
        st.caption(_t("sensitivity_no_data"))

    st.info(_t("main_info"))

    st.caption(_t("sensitivity_reminder"))
else:
    st.caption(_t( "select_at_least_one_rate" ))

# ==================== 匯出 Excel ====================
st.divider()
st.subheader("📤 " + _t( "export_subheader" ))

if st.button(_t("export_button")):
    # 防禦式取得變數，避免 Streamlit rerun 時的 scope 問題
    _current_age = current_age if 'current_age' in locals() else 49
    _current_year = current_year if 'current_year' in locals() else 2026
    _end_age = end_age if 'end_age' in locals() else 100
    _buffer_years = buffer_years if 'buffer_years' in locals() else 1.5
    _selected_rates = selected_rates if 'selected_rates' in locals() else [0.03, 0.035, 0.04]
    _retirement_years = retirement_years if 'retirement_years' in locals() else [2028]
    _results = results if 'results' in locals() else {}
    _edited_spending = edited_spending if 'edited_spending' in locals() else spending_df if 'spending_df' in locals() else None
    _main_ret = main_ret if 'main_ret' in locals() else (_retirement_years[0] if _retirement_years else 2028)
    _spending_dict = spending_dict if 'spending_dict' in locals() else {}
    _end_year = end_year if 'end_year' in locals() else _current_year + (_end_age - _current_age)
    _planning_end_year = current_year + (end_age - current_age) if 'current_year' in locals() and 'end_age' in locals() else 2080
    _living_name = living_name if 'living_name' in locals() else "Living Cost"
    _edu1_name = edu1_name if 'edu1_name' in locals() else "Kate 學費"
    _edu2_name = edu2_name if 'edu2_name' in locals() else "Damon 學費"
    _mort_name = mort_name if 'mort_name' in locals() else "Mortgage"
    _other_name = other_name if 'other_name' in locals() else "Other"
    _living_col = living_col if 'living_col' in locals() else f"{_living_name} (HKD)"
    _edu1_col = edu1_col if 'edu1_col' in locals() else f"{_edu1_name} (HKD)"
    _edu2_col = edu2_col if 'edu2_col' in locals() else f"{_edu2_name} (HKD)"
    _mort_col = mort_col if 'mort_col' in locals() else f"{_mort_name} (HKD)"
    _other_col = other_col if 'other_col' in locals() else f"{_other_name} (HKD)"
    _living_display = _living_name
    _edu1_display = _edu1_name
    _edu2_display = _edu2_name
    _mort_display = _mort_name
    _other_display = _other_name

    output = BytesIO()
    wb = Workbook()

    # ========== Sheet 1: Summary (完整參數 + 所需本金) ==========
    ws = wb.active
    ws.title = "Summary"

    ws["A1"] = _t("export_summary_title")
    ws["A1"].font = Font(bold=True, size=16, color="1F4E79")
    ws.merge_cells('A1:F1')

    ws["A3"] = _t("basic_params")
    ws["A3"].font = Font(bold=True, size=12)
    ws["A4"] = f"{_t('current_age_label')}: {_current_age}"
    ws["A5"] = f"{_t('current_year_label')}: {_current_year}"
    ws["A6"] = f"{_t('end_age_label')}: {_end_age}"
    ws["A7"] = f"{_t('buffer_label')}: {_buffer_years}"
    ws["A9"] = _t("selected_rates_label")
    for i, r in enumerate(_selected_rates):
        ws.cell(row=10+i, column=1, value=f"{r*100:.1f}%")
    ws["A15"] = _t("retirement_years_label")
    for i, ry in enumerate(_retirement_years):
        ws.cell(row=16+i, column=1, value=ry)

    # 所需本金表格 - 清楚標示每個利率
    row = 20
    ws[f"A{row}"] = _t("required_capital_title")
    ws[f"A{row}"].font = Font(bold=True, size=12)
    row += 1
    ws.cell(row=row, column=1, value=_t("retirement_year_col")).font = Font(bold=True)
    for j, r in enumerate(_selected_rates):
        pct = f"{r*100:.1f}%"
        cell = ws.cell(row=row, column=2+j, value=pct + (" Required Capital" if lang_code == "en" else " 所需本金"))
        cell.font = Font(bold=True)
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.font = Font(bold=True, color="FFFFFF")
    for i, ret_y in enumerate(_retirement_years):
        ws.cell(row=row+1+i, column=1, value=ret_y)
        for j, r in enumerate(_selected_rates):
            val = _results.get(ret_y, {}).get(r, 0)
            cell = ws.cell(row=row+1+i, column=2+j, value=val)
            cell.number_format = '#,##0'

    ws.cell(row=row+1+len(_retirement_years)+1, column=1, value=_t("note_projections"))

    # ========== Sheet 2: Full_Yearly_Spending (完整所有年份) ==========
    ws2 = wb.create_sheet("Full_Yearly_Spending")

    ws2["A1"] = f"{_t('full_spending_title')} (2025-{_planning_end_year})"
    ws2["A1"].font = Font(bold=True, size=14, color="1F4E79")
    ws2.merge_cells('A1:G1')

    ws2["A2"] = f"{_t('export_time')}: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
    ws2["A3"] = f"{_t('full_spending_note')}"

    headers = [_t("retirement_year_col"), f"{_living_display} (HKD)", f"{_edu1_display} (HKD)", f"{_edu2_display} (HKD)", f"{_mort_display} (HKD)", f"{_other_display} (HKD)", "Total (HKD)"]
    for c, h in enumerate(headers, 1):
        cell = ws2.cell(row=5, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Full data from edited_spending
    _edited = _edited_spending if _edited_spending is not None else (spending_df if 'spending_df' in locals() else None)
    if _edited is not None:
        for i, row in _edited.iterrows():
            r = 6 + i
            yr = int(row["Year"])
            liv = int(row[_living_col])
            e1  = int(row[_edu1_col])
            e2  = int(row[_edu2_col])
            mor = int(row[_mort_col])
            oth = int(row[_other_col])
            tot = liv + e1 + e2 + mor + oth

            ws2.cell(row=r, column=1, value=yr)
            ws2.cell(row=r, column=2, value=liv).number_format = '#,##0'
            ws2.cell(row=r, column=3, value=e1).number_format = '#,##0'
            ws2.cell(row=r, column=4, value=e2).number_format = '#,##0'
            ws2.cell(row=r, column=5, value=mor).number_format = '#,##0'
            ws2.cell(row=r, column=6, value=oth).number_format = '#,##0'
            ws2.cell(row=r, column=7, value=tot).number_format = '#,##0'

            for c in range(1, 8):
                ws2.cell(row=r, column=c).border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))

    # Column widths
    ws2.column_dimensions['A'].width = 10
    for col in range(2, 8):
        ws2.column_dimensions[get_column_letter(col)].width = 18

    # ========== Detailed sheets for EVERY retirement year (完整所有情境的走勢) ==========
    # 現在匯出會為你列表中的每個退休年份都產生 Projections_ (複利模型) + Simple_Model_ (像圖片的簡單模型)
    # 這樣 export 真正包含「所有 data」，不用依賴畫面上選擇的主要退休年份
    _sp_dict = _spending_dict or {}
    for ret_y in _retirement_years:
        if ret_y not in _results:
            continue

        # --- Compound Projections (用來求所需本金的模擬) ---
        ws_proj = wb.create_sheet(f"Projections_{ret_y}")
        proj_title = f"Detailed Capital Projections - Retirement Year {ret_y} (Compound Model)" if lang_code == "en" else f"資本走勢詳細表 - 退休年 {ret_y} (複利模型)"
        ws_proj["A1"] = proj_title
        ws_proj["A1"].font = Font(bold=True, size=14)

        headers3 = ["Year", "Expense (HKD)"]
        for r in _selected_rates:
            headers3 += [f"Capital_{r*100:.1f}% (HKD)"]

        for c, h in enumerate(headers3, 1):
            cell = ws_proj.cell(row=3, column=c, value=h)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        all_projections = {}
        years_list = None
        for r in _selected_rates:
            req = _results[ret_y][r]
            sim = simulate(req, r, ret_y, _end_year, _sp_dict, _buffer_years)
            all_projections[r] = sim["balances"]
            if years_list is None:
                years_list = sim["years"]

        if years_list is None:
            years_list = []

        for i, y in enumerate(years_list):
            r = 4 + i
            exp = _sp_dict.get(y, 0)
            ws_proj.cell(row=r, column=1, value=y)
            ws_proj.cell(row=r, column=2, value=exp).number_format = '#,##0'
            for j, rate in enumerate(_selected_rates):
                bal = all_projections[rate][i] if i < len(all_projections.get(rate, [])) else None
                cell = ws_proj.cell(row=r, column=3+j, value=bal if bal is not None else "")
                cell.number_format = '#,##0'

        for col in range(1, 3 + len(_selected_rates)):
            ws_proj.column_dimensions[get_column_letter(col)].width = 20

        # --- Simple Model (你圖片中的格式：Capital + Passive=cap*rate + Expense + Gain/Loss) ---
        if _selected_rates:
            ws_simple = wb.create_sheet(f"Simple_Model_{ret_y}")

            ws_simple["A1"] = (f"Capital Trajectory Table (Simple Model) - Retirement Year {ret_y} (Passive = Capital × Rate, Gain/Loss added to next Capital, matching your image format)" if lang_code == "en" else f"資本走勢表 (Simple Model) - 退休年 {ret_y} (Passive = Capital × Rate, Gain/Loss 直接加到下年 Capital，像你圖片格式)")
            ws_simple["A1"].font = Font(bold=True, size=14, color="1F4E79")
            ws_simple.merge_cells('A1:H1')
            ws_simple["A2"] = ("This table uses the Simple Model (NOT the compound simulation used to solve for required capital). All Expense values come from the yearly spending table you edited. Includes all selected rates." if lang_code == "en" else "此表使用簡單模型（非用於求解所需本金的複利模擬）。所有 Expense 來自你編輯的年度開支表。包含所有選擇的回報率。")
            ws_simple["A3"] = f"{_t( 'export_time' )}: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"

            simple_data_by_rate = {}
            base_cy = _current_year
            base_ca = _current_age
            for r in _selected_rates:
                req = _results.get(ret_y, {}).get(r, 0)
                rows = []
                capital = float(req)
                for idx, y in enumerate(range(ret_y, _end_year + 1)):
                    age = base_ca + (y - base_cy)
                    expense = _sp_dict.get(y, 0)
                    passive = capital * r
                    gain_loss = passive - expense
                    next_capital = capital + gain_loss
                    rows.append({
                        "No.": idx + 1,
                        "Age": age,
                        "YEAR": y,
                        "Expense": expense,
                        "Capital": round(capital),
                        "Passive": round(passive),
                        "Gain_Loss": round(gain_loss),
                    })
                    capital = next_capital
                simple_data_by_rate[r] = rows

            current_row = 5
            hdr = ["No.", "Age", "YEAR", "Expense (HKD)"]
            for r in _selected_rates:
                pct = f"{r*100:.1f}%"
                hdr += [f"Capital {pct}", f"Passive {pct}", f"Gain/Loss {pct}"]
            for c, h in enumerate(hdr, 1):
                cell = ws_simple.cell(row=current_row, column=c, value=h)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
                cell.border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                cell.alignment = Alignment(horizontal='center', wrap_text=True)

            current_row += 1

            n_years = len(next(iter(simple_data_by_rate.values()))) if simple_data_by_rate else 0
            for i in range(n_years):
                first_r = _selected_rates[0]
                base = simple_data_by_rate[first_r][i]
                ws_simple.cell(row=current_row, column=1, value=base["No."])
                ws_simple.cell(row=current_row, column=2, value=base["Age"])
                ws_simple.cell(row=current_row, column=3, value=base["YEAR"])
                exp_cell = ws_simple.cell(row=current_row, column=4, value=base["Expense"])
                exp_cell.number_format = '#,##0'
                col = 5
                for r in _selected_rates:
                    d = simple_data_by_rate[r][i]
                    for k in ["Capital", "Passive", "Gain_Loss"]:
                        cell = ws_simple.cell(row=current_row, column=col, value=d[k])
                        cell.number_format = '#,##0'
                        col += 1
                for c in range(1, col):
                    ws_simple.cell(row=current_row, column=c).border = Border(left=Side(style='thin'), right=Side(style='thin'), top=Side(style='thin'), bottom=Side(style='thin'))
                current_row += 1

            # 欄寬
            ws_simple.column_dimensions['A'].width = 6
            ws_simple.column_dimensions['B'].width = 6
            ws_simple.column_dimensions['C'].width = 8
            ws_simple.column_dimensions['D'].width = 15
            for c in range(5, 4 + 3 * len(_selected_rates) + 1):
                ws_simple.column_dimensions[get_column_letter(c)].width = 16

    wb.save(output)
    output.seek(0)

    st.download_button(
        label=_t("download_label"),
        data=output,
        file_name=f"{_t('export_filename')}_{datetime.datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )

st.caption(_t( "tip_caption" ))

# 免責
st.divider()
st.caption(_t("disclaimer"))

st.divider()
st.markdown(
    "**想讓其他人試用？** "
    "這個 App 可以輕鬆部署到網路上（免費）。"
    "詳細步驟請看同資料夾的 `README.md`（推薦使用 Streamlit Community Cloud 或 Hugging Face Spaces）。"
    "把原始碼 push 到 GitHub 後，幾分鐘就能有公開網址。"
)
st.caption("原始碼開源，歡迎 fork 後改成你自己的退休規劃工具。")
